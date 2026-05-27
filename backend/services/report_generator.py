from __future__ import annotations

import tempfile
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

from ..models import InvestorInput, RankedRegion, Region, RegionBrief


REPORT_DIR = Path(tempfile.gettempdir()) / "naslediye_industrii_reports"
REPORT_DIR.mkdir(parents=True, exist_ok=True)

TNR = "Times New Roman"

STATUS_LABELS = {
    "official_verified": "официально подтверждено",
    "official_stat": "официальная статистика",
    "imported_official": "импортировано из официальной выгрузки",
    "partially_verified": "частично проверено",
    "official_name_demo_metrics": "объект подтверждён, метрики оценочные",
    "demo_estimate": "демонстрационная оценка MVP",
    "needs_review": "требует проверки",
    "demo": "демонстрационная оценка MVP",
    "registry_verified": "сверено с реестром",
    "calculated_from_verified": "рассчитано автоматически",
    "data_confirmed": "автоматически сверено",
    "real_object": "объект найден",
    "tariff_estimate": "тарифная оценка",
    "market_stat_estimate": "рыночная оценка",
}

FIELD_LABELS = {
    "site_identity": "Идентичность площадки",
    "coordinates": "Координаты",
    "oez_tor_status": "Статус ОЭЗ/ТОР",
    "tax_benefits": "Налоговые льготы",
    "urban_env_index": "Индекс городской среды",
    "salary": "Средняя заработная плата",
    "kindergarten": "Детские сады",
    "profile_college": "Профильные колледжи",
    "rent": "Аренда жилья",
    "gas": "Газоснабжение",
    "power_kva": "Свободная мощность",
    "free_power_kva": "Свободная мощность",
    "substation_distance": "Расстояние до подстанции",
    "connection_fee": "Техприсоединение",
    "steel_supplier": "Поставщик стали",
    "steel_supplier_distance": "Расстояние до поставщика стали",
    "insulation_supplier": "Поставщик утеплителя",
    "insulation_supplier_distance": "Расстояние до поставщика утеплителя",
    "highway": "Федеральная трасса",
    "railway": "Ж/д доступ",
}

ARCH_LABELS = {
    "authenticity": "аутентичность региону",
    "techno": "техно-стиль",
    "eco": "экодизайн",
}

HOUSING_LABELS = {
    "dormitory": "общежитие",
    "apartment": "квартиры",
}

AMENITY_LABELS = {
    "alley": "аллея",
    "square": "сквер с фонтаном",
    "gazebo": "беседки",
    "stage": "сцена",
    "health_trail": "тропа здоровья",
    "pond": "пруд",
    "art_object": "арт-объект",
}

SPORT_LABELS = {
    "outdoor_gym": "уличные тренажёры",
    "stadium": "стадион",
    "pool": "бассейн",
    "gym": "спортзал",
    "hockey_rink": "хоккейная коробка",
}


def _safe_name(text: str) -> str:
    allowed = "abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789_-"
    out = []
    for ch in text:
        if ch in allowed:
            out.append(ch)
        elif ch.isspace() or ch in "«»\"'.,():/\\":
            out.append("_")
    return "".join(out).strip("_")[:80] or "report"


def _font(size: int = 22, bold: bool = False):
    candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSerif-Bold.ttf" if bold else "/usr/share/fonts/truetype/liberation2/LiberationSerif-Regular.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size=size)
    return ImageFont.load_default()


def _value(obj):
    return obj.value if hasattr(obj, "value") else str(obj)


def _status(value: str | None) -> str:
    return STATUS_LABELS.get(value or "", value or "не указан")


def _label_list(values, mapping: dict[str, str]) -> str:
    vals = [_value(v) for v in values or []]
    return ", ".join(mapping.get(v, v) for v in vals) if vals else "не выбрано"


def _set_run_font(run, size: float = 14, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = TNR
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run._element.rPr.rFonts.set(qn("w:eastAsia"), TNR)


def _set_paragraph_format(p, first_line: bool = True) -> None:
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    if first_line:
        p.paragraph_format.first_line_indent = Cm(1.25)


def _setup_document(doc: Document) -> None:
    section = doc.sections[0]
    # Часто используемая академическая схема оформления: левое 30 мм, правое 10 мм, верх/низ 20 мм.
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = TNR
    normal.font.size = Pt(14)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), TNR)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)

    for name in ["Heading 1", "Heading 2", "Heading 3"]:
        st = styles[name]
        st.font.name = TNR
        st.font.size = Pt(14)
        st.font.bold = True
        st._element.rPr.rFonts.set(qn("w:eastAsia"), TNR)
        st.paragraph_format.line_spacing = 1.5
        st.paragraph_format.space_before = Pt(12)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.first_line_indent = Cm(0)

    # Номер страницы внизу справа.
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    _set_run_font(run, 14)


def _add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    _set_run_font(run, 14, bold=True)


def _add_para(doc: Document, text: str, *, bold_prefix: str | None = None, italic: bool = False) -> None:
    p = doc.add_paragraph()
    _set_paragraph_format(p, first_line=True)
    text = str(text or "")
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        _set_run_font(r1, 14, bold=True, italic=italic)
        r2 = p.add_run(text[len(bold_prefix):])
        _set_run_font(r2, 14, italic=italic)
    else:
        r = p.add_run(text)
        _set_run_font(r, 14, italic=italic)


def _add_no_indent_para(doc: Document, text: str, *, bold: bool = False, align=None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Cm(0)
    if align is not None:
        p.alignment = align
    r = p.add_run(str(text))
    _set_run_font(r, 14, bold=bold)


def _add_bullets(doc: Document, items: Iterable[str]) -> None:
    items = [str(x).strip() for x in items if str(x).strip()]
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(1.25)
        r = p.add_run(item)
        _set_run_font(r, 14)


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_cell_text(cell, text: str, bold: bool = False, size: float = 12) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(str(text))
    _set_run_font(run, size, bold=bold)


def _add_table(doc: Document, headers: Sequence[str], rows: Sequence[Sequence[str]], *, font_size: float = 12) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        _set_cell_text(hdr[i], h, True, size=font_size)
        _set_cell_shading(hdr[i], "E9DED2")
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            _set_cell_text(cells[i], val, False, size=font_size)
    doc.add_paragraph()


def _add_key_value_table(doc: Document, rows, *, font_size: float = 12) -> None:
    _add_table(doc, ["Показатель", "Значение"], rows, font_size=font_size)


def _project_rows(inp: InvestorInput):
    return [
        ("Объём выпуска", f"{inp.production_volume_kt} тыс. м² панелей/год"),
        ("Количество сотрудников", f"{inp.employees} чел."),
        ("Бюджет на участок и сети", f"{inp.budget_mln_rub} млн руб."),
        ("Необходимость ж/д ветки", "требуется" if inp.needs_railway else "не требуется"),
        ("Максимальное расстояние до федеральной трассы", f"{inp.max_highway_km} км"),
        ("Архитектурный приоритет", ARCH_LABELS.get(_value(inp.arch_priority), _value(inp.arch_priority))),
        ("Обеспечение жильём", f"{inp.housing.pct}% сотрудников; тип: {HOUSING_LABELS.get(_value(inp.housing.type), _value(inp.housing.type))}"),
        ("Детский сад", f"{inp.kindergarten_per_100} мест на 100 сотрудников"),
        ("Благоустройство", _label_list(inp.amenities, AMENITY_LABELS)),
        ("Спортобъекты", _label_list(inp.sport_objects, SPORT_LABELS)),
    ]


def _region_fact_rows(region: Region, ranked: RankedRegion):
    return [
        ("Площадка", region.name),
        ("Субъект РФ", region.federal_subject),
        ("Координаты", f"{region.center.lat:.6f}, {region.center.lon:.6f}"),
        ("Итоговый рейтинг", f"{ranked.score.final_score:.3f}"),
        ("Расстояние до федеральной трассы", f"{region.logistics.federal_highway_km} км"),
        ("Ж/д доступ", "есть" if region.logistics.railway_available else "нет"),
        ("Свободная мощность", f"{region.infrastructure.free_power_kva} кВА"),
        ("Расчётная потребность", f"{ranked.required_power_kva} кВА"),
        ("Газ", "есть" if region.infrastructure.gas_available else "нет"),
        ("Подстанция", f"{region.infrastructure.substation_distance_km} км"),
        ("Плата за техприсоединение", f"{region.infrastructure.connection_fee_rub_kw} руб./кВт"),
        ("Поставщик стали", f"{region.logistics.steel_supplier_km} км"),
        ("Поставщик утеплителя", f"{region.logistics.insulation_supplier_km} км"),
        ("Энерготариф", f"{region.economy.energy_tariff_rub_kwh} руб./кВт·ч"),
        ("Средняя зарплата", f"{region.economy.avg_salary_rub} руб./мес."),
        ("Аренда 1-комнатной квартиры", f"{region.social.rent_1room_rub} руб./мес."),
        ("Индекс городской среды", f"{region.social.urban_env_index} баллов"),
        ("Обеспеченность детсадами", f"{region.social.kindergarten_per_100} мест на 100 детей"),
        ("Профильный колледж", "есть" if region.social.has_profile_college else "нет"),
        ("Экологический класс", region.economy.ecology_class),
    ]


def _score_rows(ranked: RankedRegion):
    return [
        ("Итоговый балл", f"{ranked.score.final_score:.3f}"),
        ("Логистика", f"{ranked.score.logistics:.2f}"),
        ("Экономика", f"{ranked.score.economy:.2f}"),
        ("Сетевая инфраструктура", f"{ranked.score.infrastructure:.2f}"),
        ("Социальный блок", f"{ranked.score.social:.2f}"),
        ("Бонусы", "; ".join(ranked.score.bonuses or ["нет"])),
    ]


def _area_rows(ranked: RankedRegion):
    a = ranked.areas
    return [
        ("Цех", f"{a.workshop_m2:.0f} м²"),
        ("Склад", f"{a.warehouse_m2:.0f} м²"),
        ("АБК", f"{a.office_m2:.0f} м²"),
        ("Парковка", f"{a.parking_m2:.0f} м²"),
        ("Дороги", f"{a.roads_m2:.0f} м²"),
        ("Жильё", f"{a.housing_m2:.0f} м²"),
        ("Детский сад", f"{a.kindergarten_m2:.0f} м²"),
        ("Столовая", f"{a.canteen_m2:.0f} м²"),
        ("Медпункт", f"{a.medical_m2:.0f} м²"),
        ("Итого", f"{a.total_m2:.0f} м²"),
    ]


def _estimate_rows(ranked: RankedRegion, inp: InvestorInput):
    e = ranked.estimate
    return [
        ("Цех, склад и АБК", f"{e.construction_mln:.1f} млн руб."),
        ("Жильё", f"{e.housing_mln:.1f} млн руб."),
        ("Социальные объекты", f"{e.social_mln:.1f} млн руб."),
        ("Дороги и парковка", f"{e.infrastructure_mln:.1f} млн руб."),
        ("Благоустройство и спорт", f"{e.amenities_mln:.1f} млн руб."),
        ("Полная укрупнённая смета", f"{e.total_mln:.1f} млн руб."),
        ("Участок/сети/благоустройство", f"{ranked.site_and_network_mln:.1f} млн руб. из {inp.budget_mln_rub} млн руб."),
    ]


def _generate_risks(region: Region, ranked: RankedRegion, inp: InvestorInput) -> list[str]:
    risks = list(ranked.cons or [])
    q = region.data_quality.quality_by_field or {}
    if region.infrastructure.free_power_kva < ranked.required_power_kva:
        risks.append(f"Свободная мощность {region.infrastructure.free_power_kva} кВА ниже расчётной потребности {ranked.required_power_kva} кВА.")
    if region.infrastructure.free_power_kva - ranked.required_power_kva < 150:
        risks.append("Запас свободной электрической мощности небольшой; перед проектированием нужно получить технические условия у сетевой организации.")
    if ranked.site_and_network_mln > inp.budget_mln_rub:
        risks.append(f"Оценка затрат на участок/сети {ranked.site_and_network_mln:.1f} млн руб. превышает заданный бюджет {inp.budget_mln_rub} млн руб.")
    if region.logistics.federal_highway_km > inp.max_highway_km:
        risks.append(f"Расстояние до федеральной трассы {region.logistics.federal_highway_km} км превышает ограничение инвестора {inp.max_highway_km} км.")
    if inp.needs_railway and not region.logistics.railway_available:
        risks.append("Инвестор указал необходимость ж/д ветки, но для площадки ж/д доступ не подтверждён.")
    if not region.infrastructure.gas_available:
        risks.append("Газоснабжение не подтверждено; отопление и технологические нужды могут увеличить эксплуатационные затраты.")
    if region.social.rent_1room_rub >= 18000 and inp.housing.pct < 50:
        risks.append("Аренда жилья относительно высокая; при небольшом корпоративном жилье возможны сложности с привлечением персонала.")
    needs_review = [FIELD_LABELS.get(k, k) for k, v in q.items() if v == "needs_review"]
    if needs_review:
        risks.append("Часть данных требует сверки: " + ", ".join(needs_review[:7]) + (" и др." if len(needs_review) > 7 else "") + ".")
    demo = [FIELD_LABELS.get(k, k) for k, v in q.items() if v in {"demo_estimate", "official_name_demo_metrics"}]
    if demo:
        risks.append("Часть значений используется как оценка MVP и должна быть заменена официальными данными: " + ", ".join(demo[:6]) + (" и др." if len(demo) > 6 else "") + ".")

    # Убираем дубли, но сохраняем порядок.
    unique = []
    for item in risks:
        if item and item not in unique:
            unique.append(item)
    return unique or ["Существенные ограничения не выявлены; перед инвестиционным решением требуется стандартная инженерная и правовая проверка площадки."]


def _data_quality_rows(region: Region):
    q = region.data_quality.quality_by_field or {}
    rows = [(FIELD_LABELS.get(k, k), _status(v)) for k, v in q.items()]
    if not rows:
        rows = [("Общий статус", _status(region.data_quality.status))]
    return rows


def _draw_building(draw, x, y, w, h, d, fill, outline, label, view_shift=1):
    draw.rectangle([x, y, x + w, y + h], fill=fill, outline=outline, width=2)
    pts = [(x + w, y), (x + w + d * view_shift, y - d), (x + w + d * view_shift, y + h - d), (x + w, y + h)]
    draw.polygon(pts, fill=tuple(max(0, int(c * 0.82)) for c in fill), outline=outline)
    roof = [(x, y), (x + d * view_shift, y - d), (x + w + d * view_shift, y - d), (x + w, y)]
    draw.polygon(roof, fill=tuple(min(255, int(c * 1.08)) for c in fill), outline=outline)
    font = _font(20, True)
    draw.text((x + 8, y + 8), label, font=font, fill=(35, 35, 35))
    for i in range(max(2, int(w // 90))):
        gx = x + 18 + i * 72
        draw.rectangle([gx, y + h - 55, gx + 42, y + h - 8], fill=(55, 65, 72), outline=(20, 24, 28), width=2)


def _make_view_image(path: Path, ranked: RankedRegion, region: Region, inp: InvestorInput, view: str) -> None:
    W, H = 1200, 740
    img = Image.new("RGB", (W, H), (244, 240, 232))
    draw = ImageDraw.Draw(img)
    title_font = _font(34, True)
    label_font = _font(18, True)
    small_font = _font(15, False)

    draw.rectangle([45, 105, W - 45, H - 55], fill=(225, 219, 205), outline=(110, 101, 90), width=3)
    draw.rectangle([70, 555, W - 85, 610], fill=(105, 105, 100))
    draw.rectangle([850, 130, 930, 610], fill=(105, 105, 100))

    palette = region.culture.color_palette or ["#d8c7a7", "#8b5a2b"]

    def parse_hex(h, default=(210, 198, 180)):
        h = str(h).lstrip("#")
        try:
            return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))
        except Exception:
            return default

    base = parse_hex(palette[0])
    accent = parse_hex(palette[1], (120, 88, 55))
    tech = (198, 210, 218) if _value(inp.arch_priority) == "techno" else base
    view_shift = -1 if view in ("Север", "Запад") else 1

    _draw_building(draw, 140, 260, 380, 170, 55, tech, (75, 65, 55), "Цех", view_shift)
    _draw_building(draw, 550, 300, 245, 115, 42, (198, 192, 181), (75, 65, 55), "Склад", view_shift)
    _draw_building(draw, 740, 185, 150, 95, 35, accent, (75, 65, 55), "АБК", view_shift)

    draw.rectangle([115, 455, 315, 535], fill=(188, 184, 176), outline=(70, 70, 65), width=2)
    draw.text((128, 462), "Инженерная зона", font=label_font, fill=(45, 45, 42))
    for i in range(4):
        draw.ellipse([130 + i * 40, 492, 155 + i * 40, 530], fill=(126, 137, 145), outline=(65, 70, 75), width=2)

    draw.rectangle([340, 465, 520, 535], fill=(160, 158, 150), outline=(80, 80, 75), width=2)
    draw.text((350, 472), "Рулоны стали", font=small_font, fill=(40, 40, 40))
    for i in range(5):
        draw.ellipse([355 + i * 30, 498, 378 + i * 30, 523], fill=(150, 162, 170), outline=(80, 90, 95), width=2)
    draw.rectangle([600, 455, 820, 530], fill=(164, 163, 155), outline=(80, 80, 75), width=2)
    draw.text((612, 462), "Готовые панели", font=small_font, fill=(40, 40, 40))
    for i in range(4):
        draw.rectangle([620 + i * 42, 490, 655 + i * 42, 508], fill=(230, 224, 211), outline=(120, 120, 115))
        draw.rectangle([620 + i * 42, 510, 655 + i * 42, 526], fill=(199, 211, 216), outline=(120, 120, 115))

    if ranked.areas.housing_m2 > 0:
        _draw_building(draw, 115, 145, 135, 75, 25, (196, 177, 152), (75, 65, 55), "Жильё", view_shift)
    if ranked.areas.kindergarten_m2 > 0:
        _draw_building(draw, 285, 150, 110, 62, 20, (236, 201, 120), (75, 65, 55), "Детсад", view_shift)

    draw.rectangle([960, 155, 1110, 235], fill=(120, 164, 92), outline=(80, 120, 70), width=2)
    draw.text((978, 183), "спорт", font=label_font, fill=(255, 255, 255))
    draw.ellipse([950, 295, 1095, 385], fill=(108, 170, 190), outline=(75, 120, 140), width=2)
    draw.text((990, 330), "пруд/сквер", font=label_font, fill=(255, 255, 255))

    for i in range(4):
        x = 735 + i * 55
        draw.rectangle([x, 565, x + 35, 588], fill=(222, 226, 230), outline=(60, 60, 60), width=2)
        draw.ellipse([x + 4, 584, x + 12, 592], fill=(30, 30, 30))
        draw.ellipse([x + 24, 584, x + 32, 592], fill=(30, 30, 30))

    if inp.needs_railway:
        draw.line([60, 650, W - 70, 650], fill=(35, 35, 35), width=5)
        draw.line([60, 670, W - 70, 670], fill=(35, 35, 35), width=5)
        for x in range(70, W - 70, 55):
            draw.line([x, 640, x + 25, 680], fill=(105, 76, 50), width=4)
        draw.text((78, 616), "ж/д примыкание", font=small_font, fill=(45, 45, 45))

    draw.text((50, 32), f"3D-ракурс: {view}", font=title_font, fill=(38, 34, 30))
    draw.text((50, 74), f"{region.name} · модель предприятия по производству сэндвич-панелей", font=small_font, fill=(92, 84, 75))
    img.save(path, quality=92)


def _generate_view_images(tmp_dir: Path, ranked: RankedRegion, region: Region, inp: InvestorInput) -> list[Path]:
    paths = []
    for view in ["Юг", "Север", "Запад", "Восток"]:
        p = tmp_dir / f"view_{view}.jpg"
        _make_view_image(p, ranked, region, inp, view)
        paths.append(p)
    return paths


def _add_image_grid(doc: Document, image_paths: list[Path]) -> None:
    captions = ["Вид с юга", "Вид с севера", "Вид с запада", "Вид с востока"]
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, img_path in enumerate(image_paths):
        cell = table.rows[idx // 2].cells[idx % 2]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(captions[idx])
        _set_run_font(r, 12, bold=True)
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.add_run().add_picture(str(img_path), width=Inches(3.0))


def build_investment_report_docx(region: Region, ranked: RankedRegion, inp: InvestorInput, brief: RegionBrief) -> Path:
    out = REPORT_DIR / f"report_{_safe_name(region.id)}.docx"
    tmp_dir = REPORT_DIR / f"images_{_safe_name(region.id)}"
    tmp_dir.mkdir(parents=True, exist_ok=True)
    view_images = _generate_view_images(tmp_dir, ranked, region, inp)

    doc = Document()
    _setup_document(doc)

    _add_no_indent_para(doc, "ИНВЕСТИЦИОННЫЙ ОТЧЁТ", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _add_no_indent_para(doc, "по подбору площадки для производства сэндвич-панелей", align=WD_ALIGN_PARAGRAPH.CENTER)
    _add_no_indent_para(doc, region.name, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _add_no_indent_para(doc, f"{region.federal_subject}; координаты: {region.center.lat:.6f}, {region.center.lon:.6f}", align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()

    _add_heading(doc, "1. Исходные параметры проекта", 1)
    _add_key_value_table(doc, _project_rows(inp), font_size=12)

    _add_heading(doc, "2. Общая характеристика выбранной площадки", 1)
    _add_para(doc, "Площадка рассматривается для размещения промышленного предприятия по производству сэндвич-панелей. Подбор выполнен на основе сопоставления логистики, экономических факторов, инженерной инфраструктуры, социальной среды и требований инвестора.")
    _add_key_value_table(doc, _region_fact_rows(region, ranked), font_size=11)

    _add_heading(doc, "3. Рейтинговая оценка", 1)
    _add_para(doc, ranked.selection_reason or "Площадка выбрана на основе интегрального рейтинга логистики, экономики, инфраструктуры и социальной среды.")
    _add_key_value_table(doc, _score_rows(ranked), font_size=12)

    _add_heading(doc, "4. Обоснование выбора площадки", 1)
    _add_para(doc, brief.site_specific_analysis or ranked.ai_object_analysis or "Площадка имеет сбалансированный профиль по ключевым факторам, необходимым для размещения производства.")

    _add_heading(doc, "5. Плюсы площадки", 1)
    _add_bullets(doc, ranked.pros or ["Площадка имеет сбалансированный профиль по ключевым параметрам: логистика, инженерные сети, экономика и социальная среда."])

    _add_heading(doc, "6. Ограничения, минусы и риски", 1)
    _add_bullets(doc, _generate_risks(region, ranked, inp))

    _add_heading(doc, "7. Расчёт площадей", 1)
    _add_para(doc, "Расчёт выполнен по нормативной логике задания: отдельно учитываются производственный цех, склад, АБК, дороги, парковка, жильё, детский сад, столовая и медпункт.")
    _add_key_value_table(doc, _area_rows(ranked), font_size=12)

    _add_heading(doc, "8. Укрупнённая смета", 1)
    _add_para(doc, "Смета является предварительной укрупнённой оценкой. Бюджет инвестора в форме относится к участку, подключению к сетям и благоустройству, а не ко всей стоимости строительства.")
    _add_key_value_table(doc, _estimate_rows(ranked, inp), font_size=12)

    _add_heading(doc, "9. Социальный паспорт и удержание персонала", 1)
    _add_para(doc, brief.social_passport or "Социальный паспорт включает показатели городской среды, обеспеченности детскими садами, кадровой базы и аренды жилья.")
    _add_para(doc, brief.hr_retention_tips or "Для удержания персонала рекомендуется предусмотреть жильё, транспортную доступность, столовую, медпункт и спортивную инфраструктуру.")
    _add_key_value_table(doc, [
        ("Индекс городской среды", f"{region.social.urban_env_index} баллов"),
        ("Детские сады", f"{region.social.kindergarten_per_100} мест на 100 детей"),
        ("Профильные колледжи", "есть" if region.social.has_profile_college else "нет"),
        ("Аренда 1-комнатной квартиры", f"{region.social.rent_1room_rub} руб./мес."),
        ("Демографический профиль", region.demographics_sample.dominant_age_group),
        ("Социальная рекомендация", region.demographics_sample.social_recommendation),
    ], font_size=11)

    _add_heading(doc, "10. Экономика, сети и логистика", 1)
    _add_para(doc, brief.economy_summary or "Экономический блок учитывает льготы, страховые взносы, энерготариф, заработную плату и экологический класс региона.")
    _add_para(doc, brief.infrastructure_summary or "Сетевой блок учитывает газ, свободную мощность, расстояние до подстанции и стоимость технологического присоединения.")
    _add_para(doc, brief.logistics_summary or "Логистический блок учитывает расстояние до поставщиков стали, утеплителя, федеральной трассы и наличие ж/д доступа.")

    _add_heading(doc, "11. Архитектурный концепт и материалы", 1)
    _add_para(doc, brief.design_concept or ranked.design_summary)
    _add_para(doc, "Материалы подобраны в бюджетной промышленной логике: сэндвич-панели, окрашенный профлист, металлокассеты, фиброцементные панели на АБК и перфорированные металлические экраны. Дорогие натуральные материалы допускаются только как локальные акценты во входной группе, навигации или малых архитектурных формах.")
    _add_key_value_table(doc, [
        ("Архитектурные стили региона", ", ".join(region.culture.dominant_styles or ["не указано"])),
        ("Реалистичные материалы", ", ".join(region.culture.traditional_materials or ["не указано"])),
        ("Цветовая палитра", ", ".join(region.culture.color_palette or ["не указано"])),
    ], font_size=11)

    _add_heading(doc, "12. Планировочная концепция", 1)
    _add_para(doc, brief.layout_concept or ranked.layout_concept or "Производственный цех и склад размещаются в логистически связанной зоне, АБК и общественные функции выносятся к главному входу, грузовые и пешеходные потоки разделяются.")

    _add_heading(doc, "13. 3D-модель предприятия", 1)
    _add_para(doc, "Ниже приведены четыре автоматически сформированных ракурса 3D-макета. Изображения отражают планировочную структуру: производственный цех, склад, АБК, инженерную зону, сырьевой двор, готовые панели, дороги, социальные и спортивные объекты.")
    _add_image_grid(doc, view_images)

    _add_heading(doc, "14. Альтернативные точки", 1)
    if region.alternative_sites:
        rows = [(s.name, s.type, f"{s.lat:.6f}, {s.lon:.6f}", s.description) for s in region.alternative_sites]
        _add_table(doc, ["Название", "Тип", "Координаты", "Описание"], rows, font_size=10)
    else:
        _add_para(doc, "Альтернативные точки для выбранной площадки в локальной базе не указаны.")

    _add_heading(doc, "15. Источники и качество данных", 1)
    _add_para(doc, region.data_quality.note or "Для каждого показателя используется отдельный статус качества данных.")
    _add_key_value_table(doc, [
        ("Общий статус данных", _status(region.data_quality.status)),
        ("Дата последней сверки", region.data_quality.last_verified or "не указана"),
        ("Группы подтверждённых источников", ", ".join(region.data_quality.verified_source_groups or ["не указано"])),
    ], font_size=11)
    _add_heading(doc, "15.1. Статусы по показателям", 2)
    _add_table(doc, ["Показатель", "Статус"], _data_quality_rows(region), font_size=10)
    _add_heading(doc, "15.2. Перечень источников", 2)
    if region.data_sources:
        rows = [(s.group, s.title, s.url, s.comment or "") for s in region.data_sources]
        _add_table(doc, ["Группа", "Источник", "URL", "Комментарий"], rows, font_size=9)
    else:
        _add_para(doc, "Источники для выбранной площадки не указаны в локальной базе.")

    _add_heading(doc, "16. Вывод", 1)
    _add_para(doc, "Выбранная площадка может рассматриваться как перспективный вариант для размещения производства сэндвич-панелей при условии дополнительной проверки инженерных параметров, статуса льгот, сетевых мощностей и стоимости технологического присоединения. Итоговое инвестиционное решение должно приниматься после получения официального паспорта площадки и технических условий от профильных организаций.")

    doc.save(out)
    return out


def build_project_passport_docx(region: Region, ranked: RankedRegion, inp: InvestorInput, brief: RegionBrief) -> Path:
    """Формирует паспорт проекта по типовой структуре Минстроя.

    Структура соответствует загруженному пользователем шаблону:
    1) название проекта; 2) собственник/заказчик; 3) краткое описание;
    4) инвестиционная стоимость; 5) ключевые технологические параметры;
    6) история развития; 7) технологические решения; 8) схема реализации
    и финансирования; 9) эффект; 10) документы.
    """
    out = REPORT_DIR / f"passport_{_safe_name(region.id)}.docx"
    doc = Document()
    _setup_document(doc)

    # Чёрно-белое официальное оформление.
    for style_name in ["Normal", "Heading 1", "Heading 2", "Heading 3"]:
        st = doc.styles[style_name]
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.font.name = TNR
        st._element.rPr.rFonts.set(qn("w:eastAsia"), TNR)

    def add_center(text: str, size: float = 14, bold: bool = False, spacing_after: float = 0) -> None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(spacing_after)
        r = p.add_run(text)
        _set_run_font(r, size, bold=bold)
        r.font.color.rgb = RGBColor(0, 0, 0)

    def add_right(text: str, bold: bool = False) -> None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.2
        r = p.add_run(text)
        _set_run_font(r, 12, bold=bold)
        r.font.color.rgb = RGBColor(0, 0, 0)

    def add_section(num: int, title: str) -> None:
        _add_heading(doc, f"{num}. {title}", 1)

    def add_passport_para(text: str) -> None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(1.25)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(text or ""))
        _set_run_font(r, 14)
        r.font.color.rgb = RGBColor(0, 0, 0)

    def add_plain_table(headers, rows, font_size=10, widths=None) -> None:
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(headers):
            _set_cell_text(table.rows[0].cells[i], h, True, size=font_size)
        for row in rows:
            cells = table.add_row().cells
            for i, val in enumerate(row):
                _set_cell_text(cells[i], val, False, size=font_size)
        if widths:
            for row in table.rows:
                for i, width in enumerate(widths[:len(row.cells)]):
                    row.cells[i].width = Cm(width)
        doc.add_paragraph()

    def add_kv(rows, font_size=10, widths=(6.0, 10.0)) -> None:
        add_plain_table(["Показатель", "Значение"], rows, font_size=font_size, widths=list(widths))

    # ---------- Титульный лист ----------
    add_right("УТВЕРЖДАЮ", bold=True)
    add_right("______________________________")
    add_right("должность, Ф.И.О., подпись")
    add_right("«___» ______________ 20__ г.")
    for _ in range(4):
        doc.add_paragraph()

    add_center("ПАСПОРТ ПРОЕКТА", size=18, bold=True)
    add_center("Размещение производственного объекта", size=14, bold=True)
    add_center("«Предприятие по производству сэндвич-панелей»", size=14)
    add_center(region.name, size=14, bold=True)
    add_center(region.federal_subject, size=14)
    for _ in range(2):
        doc.add_paragraph()
    add_center("Типовая структура паспорта проекта", size=14, bold=True)
    add_center("предварительная редакция для рассмотрения руководителем и инвестором", size=12)
    for _ in range(6):
        doc.add_paragraph()
    add_center("Документ сформирован информационной системой «ПромКод»", size=12)
    add_center("Дата формирования: ____________________", size=12)
    doc.add_page_break()

    # 1
    add_section(1, "Название проекта")
    add_kv([
        ("Полное наименование проекта", f"Строительство производственного объекта по выпуску сэндвич-панелей на площадке {region.name}"),
        ("Краткое наименование проекта", "Производственный комплекс сэндвич-панелей"),
        ("Место реализации", f"{region.name}, {region.federal_subject}"),
        ("Тип работ", "строительство нового производственного объекта"),
        ("Тип объекта", "промышленное предприятие"),
        ("Целевая продукция", "сэндвич-панели"),
    ])

    # 2
    add_section(2, "Собственник / Заказчик")
    add_kv([
        ("Собственник / заказчик проекта", "указывается инвестором"),
        ("Инициатор проекта", "указывается инвестором"),
        ("Оператор сопровождения", "указывается органом сопровождения инвестиционного проекта"),
        ("Правообладатель площадки", "требует уточнения по официальному паспорту площадки"),
        ("Контактное лицо", "указывается при подаче проекта на рассмотрение"),
    ])

    # 3
    add_section(3, "Краткое описание проекта")
    add_passport_para(
        f"Проект предусматривает размещение промышленного объекта по производству сэндвич-панелей "
        f"проектной мощностью {inp.production_volume_kt} тыс. м² продукции в год. Расчётная численность "
        f"персонала составляет {inp.employees} человек. В состав объекта входят производственный цех, склад, "
        f"административно-бытовой корпус, грузовой двор, инженерная зона, зоны хранения сырья и готовых панелей, "
        f"а также социальные и благоустроительные элементы в зависимости от выбранного сценария."
    )
    add_passport_para(
        f"Площадка {region.name} рассматривается как вариант размещения объекта по результатам предварительного "
        f"подбора с учётом логистики, инженерной инфраструктуры, стоимости участка и сетей, социальной среды и "
        f"кадрового потенциала. Итоговый рейтинг площадки по модели подбора составляет {ranked.score.final_score:.3f}."
    )

    # 4
    add_section(4, "Инвестиционная стоимость")
    add_kv([
        ("Укрупнённая инвестиционная стоимость", f"{ranked.estimate.total_mln:.1f} млн руб."),
        ("Участок, сети и благоустройство", f"{ranked.site_and_network_mln:.1f} млн руб."),
        ("Ориентировочный бюджет инвестора", f"{inp.budget_mln_rub:.1f} млн руб."),
        ("Внебюджетное финансирование", "предполагается за счёт средств инвестора / заёмного финансирования"),
        ("Бюджетное финансирование", "не предусмотрено на стадии предварительного паспорта"),
        ("Гранты и меры поддержки", "уточняются по статусу ОЭЗ/ТОР и региональным программам"),
    ])
    add_passport_para(
        "Стоимость является предварительной укрупнённой оценкой и не заменяет проектно-сметную документацию, "
        "коммерческие предложения подрядчиков и официальные технические условия ресурсоснабжающих организаций."
    )

    # 5
    add_section(5, "Ключевые технологические параметры объекта")
    add_plain_table(
        ["№", "Параметр", "Значение", "Примечание"],
        [
            ("1", "Проектная мощность", f"{inp.production_volume_kt} тыс. м²/год", "задаётся инвестором"),
            ("2", "Расчётная численность персонала", f"{inp.employees} чел.", "предварительный показатель"),
            ("3", "Потребность в электрической мощности", f"{ranked.required_power_kva} кВА", "расчёт по модели"),
            ("4", "Свободная мощность площадки", f"{region.infrastructure.free_power_kva} кВА", "требует сверки"),
            ("5", "Федеральная трасса", f"{region.logistics.federal_highway_km} км", "логистический показатель"),
            ("6", "Ж/д доступ", "есть" if region.logistics.railway_available else "нет", "по данным базы"),
            ("7", "Итоговая площадь объектов", f"{ranked.areas.total_m2:.0f} м²", "предварительный расчёт"),
        ],
        font_size=9,
        widths=[1.0, 5.3, 4.0, 5.4],
    )
    add_kv(_area_rows(ranked), font_size=10)

    # 6
    add_section(6, "История развития проекта")
    add_plain_table(
        ["Этап", "Содержание", "Ориентировочный статус"],
        [
            ("1. Формирование инвестиционной идеи", "Определение продукции, мощности, численности и требований к площадке", "выполнено в рамках системы"),
            ("2. Предварительный подбор площадки", "Сравнение площадок по логистике, инфраструктуре, экономике и социальной среде", "выполнено"),
            ("3. Подготовка паспорта проекта", "Формирование документа для рассмотрения руководителем/инвестором", "текущий этап"),
            ("4. Проверка исходных данных", "Получение официального паспорта площадки, ТУ, сведений о правовом статусе", "требуется"),
            ("5. Проектирование", "Разработка ПСД, уточнение генплана, санитарных разрывов и инженерных решений", "последующий этап"),
            ("6. Строительство и ввод", "Строительно-монтажные работы, пусконаладка, ввод объекта в эксплуатацию", "последующий этап"),
        ],
        font_size=9,
        widths=[4.2, 8.0, 4.0],
    )

    # 7
    add_section(7, "Используемые технологические решения")
    add_passport_para(
        "Проект предусматривает организацию производственной линии по выпуску сэндвич-панелей с выделением "
        "зон приёмки сырья, производственного цеха, складирования готовой продукции и отгрузки. Планировочная "
        "схема предполагает разделение грузовых, производственных и пешеходных потоков."
    )
    add_kv([
        ("Производственный блок", "цех изготовления сэндвич-панелей"),
        ("Складской блок", "хранение сырья и готовой продукции"),
        ("Логистический блок", "грузовой двор, доки отгрузки, связь с автомобильной дорогой и при необходимости с ж/д"),
        ("Инженерный блок", "электроснабжение, вентиляция, технологические подключения"),
        ("Административно-бытовой блок", "управление, бытовые помещения, КПП"),
        ("Архитектурный подход", ARCH_LABELS.get(_value(inp.arch_priority), _value(inp.arch_priority))),
        ("Благоустройство", _label_list(inp.amenities, AMENITY_LABELS)),
    ])

    # 8
    add_section(8, "Схема реализации и финансирования проекта")
    add_plain_table(
        ["Участник", "Предполагаемая роль"],
        [
            ("Инвестор / заказчик", "финансирование проекта, принятие инвестиционного решения, постановка требований"),
            ("Проектная организация", "разработка проектной и рабочей документации"),
            ("Орган сопровождения инвестиционного проекта", "консультационное сопровождение, взаимодействие с площадкой и органами власти"),
            ("Оператор площадки / управляющая компания", "предоставление сведений о площадке, инфраструктуре и условиях размещения"),
            ("Финансовые институты", "кредитование, лизинг или иные финансовые инструменты при необходимости"),
            ("Ресурсоснабжающие организации", "выдача технических условий и подтверждение параметров подключения"),
        ],
        font_size=9,
        widths=[5.0, 11.0],
    )
    add_passport_para(
        "Предварительно предполагается смешанная схема финансирования с преобладанием внебюджетных средств. "
        "Конкретная структура финансирования, механизм возврата средств и участие финансовых институтов уточняются "
        "после подтверждения инвестиционной модели и получения официальных условий площадки."
    )

    # 9
    add_section(9, "Эффект от реализации проекта")
    add_plain_table(
        ["Направление эффекта", "Ожидаемый результат"],
        [
            ("Производственный эффект", f"создание мощности до {inp.production_volume_kt} тыс. м² сэндвич-панелей в год"),
            ("Занятость", f"создание или обеспечение занятости до {inp.employees} человек"),
            ("Инфраструктурный эффект", "освоение промышленной площадки и развитие инженерно-логистической инфраструктуры"),
            ("Социальный эффект", "наличие социального блока, жилья, детского сада и спортивных объектов в зависимости от выбранного сценария"),
            ("Бюджетный эффект", "формирование налоговой базы после запуска производства"),
            ("Градостроительный эффект", "размещение производственного объекта с учётом благоустройства и регионального архитектурного образа"),
        ],
        font_size=9,
        widths=[5.0, 11.0],
    )
    add_passport_para(
        "Эффект носит предварительный расчётный характер и подлежит уточнению после подтверждения мощности, "
        "состава оборудования, графика реализации, условий поддержки и финансовой модели проекта."
    )

    # 10
    add_section(10, "Документы")
    add_plain_table(
        ["№", "Документ / материал", "Статус"],
        [
            ("1", "Паспорт проекта", "сформирован информационной системой"),
            ("2", "Предварительная инженерная схема участка", "приложение к паспорту"),
            ("3", "Word-отчёт с аналитикой площадки", "формируется отдельно"),
            ("4", "Презентация проекта", "формируется отдельно"),
            ("5", "Официальный паспорт промышленной площадки", "требуется получить у оператора площадки"),
            ("6", "Технические условия на подключение", "требуется получить у ресурсоснабжающих организаций"),
            ("7", "Правоустанавливающие документы на земельный участок", "требуются для утверждения проекта"),
            ("8", "Проектно-сметная документация", "разрабатывается на последующих стадиях"),
        ],
        font_size=9,
        widths=[1.0, 10.5, 4.5],
    )

    _add_heading(doc, "10.1. Качество данных", 2)
    add_kv([
        ("Общий статус", _status(region.data_quality.status)),
        ("Дата последней сверки", region.data_quality.last_verified or "не указана"),
        ("Комментарий", region.data_quality.note or "не указан"),
    ], font_size=10)
    add_plain_table(["Показатель", "Статус"], _data_quality_rows(region), font_size=9, widths=[7.5, 8.0])

    # Приложение А.
    doc.add_page_break()
    add_center("Приложение А", bold=True)
    add_center("Предварительная инженерная схема участка", bold=True)
    try:
        from .render_generator import generate_site_plan
        plan_path = generate_site_plan(region, ranked, inp)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(plan_path), width=Cm(16.0))
    except Exception:
        _add_para(doc, "Схема участка не была добавлена автоматически. Она доступна во вкладке макета предприятия в интерфейсе системы.")

    # Лист согласования.
    doc.add_page_break()
    add_center("ЛИСТ СОГЛАСОВАНИЯ", bold=True)
    add_plain_table(
        ["№", "Должность / организация", "Ф.И.О.", "Подпись", "Дата"],
        [(str(i), "", "", "", "") for i in range(1, 7)],
        font_size=10,
        widths=[1.0, 5.0, 4.0, 3.0, 2.5],
    )

    doc.save(out)
    return out

