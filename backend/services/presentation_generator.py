from __future__ import annotations

import tempfile
from pathlib import Path
from typing import Dict

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.util import Inches, Pt

from ..models import InvestorInput, RankedRegion, Region, RegionBrief
from .render_generator import generate_project_renders, generate_site_plan, render_job_id

PRESENTATION_DIR = Path(tempfile.gettempdir()) / "naslediye_industrii_generated" / "presentations"
PRESENTATION_DIR.mkdir(parents=True, exist_ok=True)

ARCH_LABELS = {
    "authenticity": "Аутентичность региону",
    "techno": "Техно-стиль",
    "eco": "Экодизайн",
}
HOUSING_LABELS = {"dormitory": "общежитие", "apartment": "квартиры"}
AMENITY_LABELS = {
    "alley": "аллея", "square": "сквер с фонтаном", "gazebo": "беседки", "stage": "сцена",
    "health_trail": "тропа здоровья", "pond": "пруд", "art_object": "арт-объект",
}
SPORT_LABELS = {
    "outdoor_gym": "уличные тренажёры", "stadium": "стадион", "pool": "бассейн",
    "gym": "спортзал", "hockey_rink": "хоккейная коробка",
}


def _value(obj) -> str:
    return obj.value if hasattr(obj, "value") else str(obj)


def _hex_to_rgb(value: str, default=(31, 42, 64)):
    value = (value or "").strip().replace("#", "")
    if len(value) == 3:
        value = "".join(ch * 2 for ch in value)
    if len(value) != 6:
        return default
    try:
        return tuple(int(value[i:i+2], 16) for i in (0, 2, 4))
    except ValueError:
        return default


def _rgb(color) -> RGBColor:
    return RGBColor(*color)


def _set_text(shape, text: str, size=22, bold=False, color=(31,42,64), align=None):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    p = tf.paragraphs[0]
    p.text = str(text)
    if align is not None:
        p.alignment = align
    run = p.runs[0] if p.runs else p.add_run()
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)


def _add_text(slide, x, y, w, h, text, size=22, bold=False, color=(31,42,64), align=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    _set_text(shape, text, size, bold, color, align)
    return shape


def _add_rect(slide, x, y, w, h, fill, line=None, radius=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = _rgb(fill)
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = _rgb(line)
    return shp


def _add_title(slide, title: str, subtitle: str | None = None, accent=(180,126,54)):
    _add_text(slide, 0.55, 0.35, 11.8, 0.45, title, size=29, bold=True, color=(28,34,54))
    _add_rect(slide, 0.62, 0.9, 1.15, 0.06, accent)
    if subtitle:
        _add_text(slide, 0.55, 1.03, 11.3, 0.42, subtitle, size=14, color=(87,94,112))


def _add_card(slide, x, y, w, h, title, value, sub="", accent=(180,126,54)):
    _add_rect(slide, x, y, w, h, (247,246,242), (226,223,215), radius=True)
    _add_text(slide, x+0.15, y+0.12, w-0.3, 0.25, title, size=11, bold=True, color=accent)
    _add_text(slide, x+0.15, y+0.43, w-0.3, 0.42, value, size=22, bold=True, color=(28,34,54))
    if sub:
        _add_text(slide, x+0.15, y+0.92, w-0.3, h-0.98, sub, size=10.5, color=(83,90,110))


def _bullet_box(slide, x, y, w, h, title: str, items, accent=(180,126,54), font_size=11):
    _add_rect(slide, x, y, w, h, (249,248,244), (226,223,215), radius=True)
    _add_text(slide, x+0.18, y+0.14, w-0.36, 0.35, title, size=15, bold=True, color=(28,34,54))

    box = slide.shapes.add_textbox(Inches(x+0.24), Inches(y+0.56), Inches(w-0.42), Inches(h-0.72))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)

    clean_items = [str(item).strip() for item in list(items) if str(item).strip()][:5]
    for idx, item in enumerate(clean_items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.bullet = True
        p.space_after = Pt(5)
        p.alignment = PP_ALIGN.LEFT
        if p.runs:
            run = p.runs[0]
        else:
            run = p.add_run()
        run.font.name = "Arial"
        run.font.size = Pt(font_size)
        run.font.bold = False
        run.font.color.rgb = _rgb((66,72,88))


def _short(text: str, limit: int = 230) -> str:
    text = " ".join(str(text or "").split())
    return text if len(text) <= limit else text[:limit-1].rsplit(" ", 1)[0] + "…"


def build_administration_presentation_pptx(region: Region, ranked: RankedRegion, inp: InvestorInput, brief: RegionBrief, layout_key: str | None = None, custom_renders: Dict[str, Path] | None = None) -> Path:
    renders = custom_renders or generate_project_renders(region, ranked, inp, brief, layout_key)
    plan_path = generate_site_plan(region, ranked, inp, layout_key)
    job = render_job_id(region.id, inp, layout_key)
    if custom_renders:
        job = f"{job}_3dshots"
    out = PRESENTATION_DIR / f"administration_package_{job}.pptx"

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    palette = region.culture.color_palette or []
    accent = _hex_to_rgb(palette[1] if len(palette)>1 else "", (184,128,56))
    dark = (28, 34, 54)
    muted = (82, 90, 108)
    light = (247, 246, 242)

    # Slide 1: title
    s = prs.slides.add_slide(blank)
    _add_rect(s, 0, 0, 13.333, 7.5, (248,247,243))
    _add_rect(s, 0, 0, 13.333, 7.5, (255,255,255), None)
    _add_rect(s, 0, 0, 13.333, 0.18, accent)
    _add_text(s, 0.75, 0.75, 11.8, 0.35, "ПромКод", size=22, bold=True, color=accent)
    _add_text(s, 0.72, 1.32, 10.9, 1.35, "Подбор площадки и визуальный концепт промышленного объекта", size=37, bold=True, color=dark)
    _add_text(s, 0.76, 2.78, 10.9, 0.45, f"{region.name} · {region.federal_subject}", size=20, color=muted)
    _add_card(s, 0.75, 4.18, 2.75, 1.32, "Производство", f"{inp.production_volume_kt} тыс. м²/год", "сэндвич-панели", accent)
    _add_card(s, 3.75, 4.18, 2.1, 1.32, "Рабочие места", f"{inp.employees}", "сотрудников", accent)
    _add_card(s, 6.1, 4.18, 2.35, 1.32, "Смета", f"{ranked.estimate.total_mln:.0f} млн ₽", "предварительно", accent)
    _add_card(s, 8.7, 4.18, 2.3, 1.32, "Рейтинг", f"{ranked.score.final_score:.3f}", "интегральный балл", accent)

    # Slide 2: parameters
    s = prs.slides.add_slide(blank)
    _add_title(s, "Параметры проекта", "Ключевые исходные данные проекта и расчётные показатели", accent)
    _add_card(s, 0.65, 1.55, 2.35, 1.15, "Объём", f"{inp.production_volume_kt}", "тыс. м² панелей в год", accent)
    _add_card(s, 3.25, 1.55, 2.35, 1.15, "Сотрудники", f"{inp.employees}", "человек", accent)
    _add_card(s, 5.85, 1.55, 2.35, 1.15, "Бюджет", f"{inp.budget_mln_rub}", "млн ₽ на участок и сети", accent)
    _add_card(s, 8.45, 1.55, 2.35, 1.15, "Ж/д ветка", "да" if inp.needs_railway else "нет", f"трасса до {inp.max_highway_km} км", accent)
    _bullet_box(s, 0.65, 3.1, 3.8, 2.85, "Архитектура и благоустройство", [
        f"приоритет: {ARCH_LABELS.get(_value(inp.arch_priority), _value(inp.arch_priority))}",
        "благоустройство: " + (", ".join(AMENITY_LABELS.get(_value(x), _value(x)) for x in inp.amenities) or "не выбрано"),
        "материалы: " + ", ".join(region.culture.traditional_materials[:3]),
        "палитра: " + ", ".join(region.culture.color_palette[:4]),
    ], accent)
    _bullet_box(s, 4.75, 3.1, 3.8, 2.85, "Социальный пакет", [
        f"жильё: {inp.housing.pct}% сотрудников",
        f"тип жилья: {HOUSING_LABELS.get(_value(inp.housing.type), _value(inp.housing.type))}",
        f"детский сад: {inp.kindergarten_per_100} мест на 100 сотрудников",
        "спорт: " + (", ".join(SPORT_LABELS.get(_value(x), _value(x)) for x in inp.sport_objects) or "не выбран"),
    ], accent)
    _bullet_box(s, 8.85, 3.1, 3.8, 2.85, "Расчётные площади", [
        f"цех: {ranked.areas.workshop_m2:.0f} м²",
        f"склад: {ranked.areas.warehouse_m2:.0f} м²",
        f"АБК: {ranked.areas.office_m2:.0f} м²",
        f"дороги и парковка: {(ranked.areas.roads_m2 + ranked.areas.parking_m2):.0f} м²",
        f"итого: {ranked.areas.total_m2:.0f} м²",
    ], accent)

    # Slide 3: renders
    s = prs.slides.add_slide(blank)
    _add_title(s, "4 визуальных рендера", "Виды площадки на основе выбранной схемы 3D-макета: юг, север, запад и восток", accent)
    positions = [(0.7,1.45), (6.95,1.45), (0.7,4.3), (6.95,4.3)]
    labels = {"south":"Юг", "north":"Север", "west":"Запад", "east":"Восток"}
    for (key, path), (x,y) in zip(renders.items(), positions):
        s.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(5.7), height=Inches(2.55))
        _add_rect(s, x, y, 1.05, 0.32, accent, None, radius=True)
        _add_text(s, x+0.07, y+0.05, 0.9, 0.22, labels.get(key,key), size=10, bold=True, color=(255,255,255), align=PP_ALIGN.CENTER)

    # Slide 4: plan
    s = prs.slides.add_slide(blank)
    _add_title(s, "План участка с социальными объектами", "Схема отделяет грузовой контур от общественного входа и выносит социалку в тихую зону", accent)
    s.shapes.add_picture(str(plan_path), Inches(0.65), Inches(1.36), width=Inches(8.2), height=Inches(4.62))
    _bullet_box(s, 9.2, 1.45, 3.35, 4.5, "Логика размещения", [
        "цех и склад формируют производственное ядро",
        "КПП и АБК расположены у общественного входа",
        "грузовой двор выведен к трассе/ж-д направлению",
        "парковка размещается до проходной",
        "детсад, спорт и зелёные зоны отделены от грузового потока",
    ], accent)

    # Slide 5: networks and compliance
    s = prs.slides.add_slide(blank)
    _add_title(s, "Нормативы и доступность сетей", "Инженерная готовность участка и базовые производственные ограничения", accent)
    _add_card(s, 0.65, 1.55, 2.55, 1.25, "Газ", "есть" if region.infrastructure.gas_available else "нет", "магистральный газ в промзоне", accent)
    _add_card(s, 3.45, 1.55, 2.55, 1.25, "Мощность", f"{region.infrastructure.free_power_kva} кВА", f"потребность: {ranked.required_power_kva} кВА", accent)
    _add_card(s, 6.25, 1.55, 2.55, 1.25, "Подстанция", f"{region.infrastructure.substation_distance_km} км", "расстояние до ближайшей", accent)
    _add_card(s, 9.05, 1.55, 2.55, 1.25, "Техприсоединение", f"{region.infrastructure.connection_fee_rub_kw:,}".replace(',', ' '), "руб/кВт", accent)
    _bullet_box(s, 0.65, 3.25, 5.75, 2.8, "Производственные нормативы", [
        "тип производства: сэндвич-панели",
        "высота цеха: 8–10 м",
        "расчётная потребность линии: 300–800 кВА",
        "склад: 35% от площади цеха",
        "дороги: 25% от цеха и склада",
    ], accent)
    _bullet_box(s, 6.8, 3.25, 5.75, 2.8, "Сетевой вывод", [
        f"магистральный газ: {'есть' if region.infrastructure.gas_available else 'нет'}",
        f"свободная мощность: {region.infrastructure.free_power_kva} кВА",
        f"запас мощности: {region.infrastructure.free_power_kva - ranked.required_power_kva} кВА",
        f"подстанция: {region.infrastructure.substation_distance_km} км",
        f"участок и сети: {ranked.site_and_network_mln:.1f} млн ₽",
    ], accent)

    # Slide 6: economy/social effect
    s = prs.slides.add_slide(blank)
    _add_title(s, "Экономика и социальные выгоды для региона", "Почему проект может быть интересен муниципалитету", accent)
    _add_card(s, 0.65, 1.45, 2.55, 1.18, "Энерготариф", f"{region.economy.energy_tariff_rub_kwh}", "руб/кВт·ч", accent)
    _add_card(s, 3.45, 1.45, 2.55, 1.18, "Средняя ЗП", f"{region.economy.avg_salary_rub:,}".replace(',', ' '), "руб/мес", accent)
    _add_card(s, 6.25, 1.45, 2.55, 1.18, "Льготы", "есть" if region.economy.has_oez_tor else "нет", "ОЭЗ/ТОР/парк", accent)
    _add_card(s, 9.05, 1.45, 2.55, 1.18, "Рабочие места", f"{inp.employees}", "новые занятости", accent)
    benefits = []
    if region.economy.has_oez_tor:
        benefits.append("есть льготный режим для инвестора")
    if region.economy.reduced_insurance:
        benefits.append("пониженные страховые взносы")
    if not benefits:
        benefits.append("доступна стандартная региональная поддержка")

    sport_line = ", ".join(SPORT_LABELS.get(_value(x), _value(x)) for x in inp.sport_objects) or "не выбран"
    demo = getattr(region, 'demographics_sample', None)
    demo_line = (
        f"целевая группа кадров: {demo.dominant_age_group}" if demo else "проект усиливает кадровую привлекательность площадки"
    )

    _bullet_box(s, 0.65, 3.05, 3.8, 2.85, "Экономика", [
        f"энерготариф: {region.economy.energy_tariff_rub_kwh} руб/кВт·ч",
        f"средняя зарплата: {region.economy.avg_salary_rub:,} руб/мес".replace(',', ' '),
        "; ".join(benefits),
        f"строительство: {ranked.estimate.construction_mln:.0f} млн ₽",
        f"общая смета: {ranked.estimate.total_mln:.0f} млн ₽",
    ], accent)
    _bullet_box(s, 4.75, 3.05, 3.8, 2.85, "Социальный эффект", [
        f"жильё для {inp.housing.pct}% сотрудников",
        f"детский сад: {inp.kindergarten_per_100} мест на 100 сотрудников",
        f"спорт: {sport_line}",
        demo_line,
        "проект повышает устойчивость занятости в районе",
    ], accent)
    _bullet_box(s, 8.85, 3.05, 3.8, 2.85, "Рекомендация", [
        f"площадка подходит под проект объёмом {inp.production_volume_kt} тыс. м²/год",
        f"инженерные условия соответствуют потребности линии {ranked.required_power_kva} кВА",
        "проект создаёт новые рабочие места и усиливает социальную инфраструктуру",
        "рекомендуется переход к детальной проработке архитектуры и сетей",
    ], accent)

    prs.save(out)
    return out
